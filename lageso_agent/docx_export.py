"""
DOCX-Export im CU-Standard für den LaGeSo-Befundbericht.
Vollständige 12-Abschnitt-Struktur.
Formatierung: Arial 11pt, Überschriften fett 12pt, Kopfzeile 13pt.
"""

import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


def _set_run_font(run, size=11, bold=False, font_name="Arial"):
    """Setzt Schriftart-Eigenschaften für einen Run."""
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_paragraph(doc, text, size=11, bold=False, alignment=None, spacing_after=6):
    """Fügt einen formatierten Absatz hinzu."""
    para = doc.add_paragraph()
    if alignment is not None:
        para.alignment = alignment
    para.paragraph_format.space_after = Pt(spacing_after)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run(text)
    _set_run_font(run, size=size, bold=bold)
    return para


def _add_fliesstext(doc, text, spacing_after=6):
    """Fügt einen Fließtext-Absatz mit Zeilenabstand 1.3 hinzu."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(spacing_after)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.line_spacing = 1.3
    run = para.add_run(text)
    _set_run_font(run, size=11)
    return para


def _add_section_header(doc, text, spacing_after=6):
    """Fügt eine Abschnittsüberschrift hinzu (fett, 12pt)."""
    return _add_paragraph(doc, text, size=12, bold=True, spacing_after=spacing_after)


def _add_lines(doc, text, bullet=False, size=11, spacing_after=2):
    """Fügt mehrzeiligen Text als einzelne Absätze hinzu."""
    for line in text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        display = f"\u2022 {line}" if bullet else line
        _add_paragraph(doc, display, size=size, spacing_after=spacing_after)


def export_docx(
    arzt_daten: dict,
    patient_name: str,
    geburtsdatum: str,
    plz_ort: str,
    strasse: str,
    aktenzeichen: str,
    vorstellung_text: str,
    diagnosen_text: str,
    befund_text: str,
    befund_datum: str,
    medikation_text: str,
    verlauf_text: str,
    sucht_text: str,
    behandlungen_text: str,
    psychosoziale_hilfen_text: str,
    auswirkungen_alltag: str,
    auswirkungen_beruf: str,
    auswirkungen_sozial: str,
    gehfaehigkeit_text: str,
    akteneinsicht_text: str,
) -> bytes:
    """Exportiert den vollständigen 12-Abschnitt-Bericht als DOCX.

    Returns:
        DOCX-Datei als bytes
    """
    doc = Document()
    heute = date.today().strftime("%d.%m.%Y")

    # --- Seitenränder ---
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # --- Kopfzeile: Arzt-Stammdaten ---
    _add_paragraph(doc, arzt_daten.get("name", ""), size=13, bold=True, spacing_after=2)
    _add_paragraph(doc, arzt_daten.get("fachgebiet", ""), size=10, spacing_after=2)
    _add_paragraph(doc, arzt_daten.get("adresse", ""), size=10, spacing_after=2)
    _add_paragraph(doc, arzt_daten.get("kontakt", ""), size=10, spacing_after=12)

    # --- Datum rechtsbündig ---
    _add_paragraph(
        doc, f"Berlin, den {heute}",
        size=11, alignment=WD_ALIGN_PARAGRAPH.RIGHT, spacing_after=12,
    )

    # --- Titel ---
    _add_paragraph(
        doc, "Anlage zur Auskunft über die ärztliche Behandlung (LaGeSo)",
        size=11, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, spacing_after=2,
    )
    _add_paragraph(
        doc, "Psychiatrischer/psychotherapeutischer Befundbericht",
        size=11, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, spacing_after=12,
    )

    # --- Patienten-Kopf ---
    adresse = f"{plz_ort}, {strasse}" if plz_ort and strasse else (plz_ort or strasse or "[Adresse]")
    _add_paragraph(
        doc,
        f"Patient/in {patient_name or '[Name]'}, "
        f"geb. am {geburtsdatum or '[Geburtsdatum]'}, "
        f"wohnhaft in {adresse}",
        size=11, spacing_after=2,
    )
    _add_paragraph(doc, f"Aktenzeichen: {aktenzeichen or '[Aktenzeichen]'}", size=11, spacing_after=12)

    # === 1. VORSTELLUNG ===
    _add_section_header(doc, "1. VORSTELLUNG")
    _add_paragraph(doc, "Die Vorstellung erfolgte:", size=11, spacing_after=2)
    _add_lines(doc, vorstellung_text)
    _add_paragraph(doc, "", spacing_after=6)

    # === 2. DIAGNOSEN ===
    _add_section_header(doc, "2. DIAGNOSEN (ICD-10)")
    for line in diagnosen_text.strip().split("\n"):
        line = line.strip()
        if line:
            _add_paragraph(doc, f"\u2022 {line}", size=11, spacing_after=2)
    _add_paragraph(doc, "", spacing_after=6)

    # === 3. PSYCHOPATHOLOGISCHER BEFUND ===
    _add_section_header(doc, "3. PSYCHOPATHOLOGISCHER BEFUND", spacing_after=2)
    _add_paragraph(
        doc,
        f"Letzter vollständiger psychopathologischer Befund \u2013 "
        f"nach dem AMDP-System (Datum: {befund_datum or heute})",
        size=10, spacing_after=6,
    )
    _add_fliesstext(doc, befund_text, spacing_after=12)

    # === 4. AKTUELLE MEDIKATION ===
    _add_section_header(doc, "4. AKTUELLE MEDIKATION (nur psychopharmakologisch)")
    for line in medikation_text.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("("):
            _add_paragraph(doc, line, size=10, spacing_after=2)
        else:
            _add_paragraph(doc, f"\u2022 {line}", size=11, spacing_after=2)
    _add_paragraph(doc, "", spacing_after=6)

    # === 5. VERLAUF DER ERKRANKUNG ===
    _add_section_header(doc, "5. VERLAUF DER ERKRANKUNG")
    _add_fliesstext(doc, verlauf_text, spacing_after=12)

    # === 6. BEI SUCHTERKRANKUNG ===
    _add_section_header(doc, "6. BEI SUCHTERKRANKUNG")
    _add_fliesstext(doc, sucht_text, spacing_after=12)

    # === 7. SONSTIGE BEHANDLUNGEN ===
    _add_section_header(doc, "7. SONSTIGE BEHANDLUNGEN (wo?, wann?)")
    _add_paragraph(
        doc, "NUR psychiatrische/psychotherapeutische Behandlungen.",
        size=10, spacing_after=4,
    )
    _add_lines(doc, behandlungen_text)
    _add_paragraph(doc, "", spacing_after=6)

    # === 8. PSYCHOSOZIALE HILFEN ===
    _add_section_header(doc, "8. PSYCHOSOZIALE HILFEN")
    _add_lines(doc, psychosoziale_hilfen_text)
    _add_paragraph(doc, "", spacing_after=6)

    # === 9. KRANKHEITSBEDINGTE AUSWIRKUNGEN ===
    _add_section_header(doc, "9. KRANKHEITSBEDINGTE AUSWIRKUNGEN")

    _add_paragraph(doc, "Alltag:", size=11, bold=True, spacing_after=2)
    _add_fliesstext(doc, auswirkungen_alltag, spacing_after=8)

    _add_paragraph(doc, "Beruf:", size=11, bold=True, spacing_after=2)
    _add_fliesstext(doc, auswirkungen_beruf, spacing_after=8)

    _add_paragraph(doc, "Sozial:", size=11, bold=True, spacing_after=2)
    _add_fliesstext(doc, auswirkungen_sozial, spacing_after=12)

    # === 10. EINSCHRÄNKUNGEN DER GEHFÄHIGKEIT ===
    _add_section_header(doc, "10. EINSCHRÄNKUNGEN DER GEHFÄHIGKEIT")
    _add_paragraph(doc, gehfaehigkeit_text, size=11, spacing_after=12)

    # === 11. AKTENEINSICHT ===
    _add_section_header(doc, "11. AKTENEINSICHT")
    _add_paragraph(
        doc,
        "Im Falle einer Akteneinsicht kann dieser Befundbericht "
        "dem/der Antragsteller/in zur Einsicht gegeben werden?",
        size=11, spacing_after=4,
    )
    _add_paragraph(doc, akteneinsicht_text, size=11, spacing_after=12)

    # === 12. ABSCHLUSS ===
    _add_section_header(doc, "12. ABSCHLUSS")
    _add_paragraph(
        doc,
        "Den Befundbericht bitte ich gemäß beiliegender Liquidation zu entschädigen.",
        size=11, spacing_after=8,
    )
    _add_paragraph(doc, "Mit freundlichen Grüßen", size=11, spacing_after=24)

    # --- Unterschrift ---
    _add_paragraph(doc, "___________________________", size=11, spacing_after=2)
    _add_paragraph(doc, arzt_daten.get("name", ""), size=11, spacing_after=0)

    # --- In Bytes konvertieren ---
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
